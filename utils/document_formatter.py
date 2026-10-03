from __future__ import annotations

import html
import re

from io import BytesIO
from pathlib import Path

from docx import Document

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF


ROOT_DIR = Path(__file__).resolve().parents[2]

LOGO_PATH = (
    ROOT_DIR /
    "assets" /
    "logo.png"
)


def sanitize_text(text: str) -> str:

    replacements = {

        "\u2018": "'",
        "\u2019": "'",

        "\u201c": '"',
        "\u201d": '"',

        "\u2013": "-",
        "\u2014": "-",

        "\u2026": "...",

        "\u00a0": " "
    }

    for old, new in replacements.items():

        text = text.replace(
            old,
            new
        )

    return "".join(
        character
        for character in text
        if (
            character in "\n\r\t"
            or ord(character) >= 32
        )
    ).strip()


def split_lines(text: str) -> list[str]:

    cleaned = sanitize_text(text)

    return [
        line.strip()
        for line in cleaned.splitlines()
    ]


def is_heading(line: str) -> bool:

    if not line:
        return False

    if re.match(
        r"^\d+[\.\)]\s+",
        line
    ):
        return True

    if re.match(
        r"^(ARTICLE|SECTION)\s+[A-Z0-9IVX]+",
        line,
        re.IGNORECASE
    ):
        return True

    if (
        len(line) <= 70
        and line.isupper()
        and len(line.split()) <= 10
    ):
        return True

    return False


def is_bullet(line: str) -> bool:

    return bool(
        re.match(
            r"^[-*•]\s+",
            line
        )
    )


def format_docx(
    text: str,
    doc_type: str,
    terms: str = ""
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles

    styles["Normal"].font.name = (
        "Times New Roman"
    )

    styles["Normal"].font.size = Pt(11)

    # Logo

    if LOGO_PATH.exists():

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.add_run().add_picture(
            str(LOGO_PATH),
            width=Inches(1.45)
        )

    # Title

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        doc_type.strip()
        or "Legal Document"
    )

    title_run.bold = True
    title_run.font.name = (
        "Times New Roman"
    )
    title_run.font.size = Pt(16)

    # Main document

    for line in split_lines(text):

        if not line:

            document.add_paragraph("")

            continue

        if is_heading(line):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(line)

            run.bold = True
            run.font.name = (
                "Times New Roman"
            )
            run.font.size = Pt(12)

        elif is_bullet(line):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                re.sub(
                    r"^[-*•]\s+",
                    "",
                    line
                )
            )

        else:

            paragraph = document.add_paragraph(
                line
            )

            paragraph.paragraph_format.space_after = (
                Pt(6)
            )

            paragraph.paragraph_format.line_spacing = (
                1.15
            )

    # Terms table

    clean_terms = [
        item.strip()
        for item in terms.split(";")
        if item.strip()
    ]

    if clean_terms:

        document.add_page_break()

        heading = document.add_paragraph()

        run = heading.add_run(
            "Key Terms Provided"
        )

        run.bold = True
        run.font.size = Pt(13)

        table = document.add_table(
            rows=1,
            cols=2
        )

        table.style = "Table Grid"

        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = (
            "Term / Condition"
        )

        for index, term in enumerate(
            clean_terms,
            start=1
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)
            cells[1].text = term

    # Footer

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase - AI-generated draft. Review before use."
    )

    footer_run.italic = True

    # Save to memory

    buffer = BytesIO()

    document.save(buffer)

    return buffer.getvalue()


class LegalEasePDF(FPDF):

    def __init__(
        self,
        doc_type: str
    ):

        super().__init__()

        self.doc_type = doc_type

        self.set_auto_page_break(
            auto=True,
            margin=18
        )

    def header(self):

        if LOGO_PATH.exists():

            try:

                self.image(
                    str(LOGO_PATH),
                    x=95,
                    y=7,
                    w=20
                )

                self.set_y(29)

            except Exception:

                self.set_y(10)

        else:

            self.set_y(10)

        self.set_font(
            "Helvetica",
            "B",
            11
        )

        self.cell(
            0,
            7,
            self.doc_type[:90],
            align="C"
        )

        self.ln(8)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "I",
            8
        )

        self.cell(
            0,
            8,
            "LegalEase - AI-generated draft. Review before use.",
            align="C"
        )


def format_pdf(
    text: str,
    doc_type: str,
    terms: str = ""
) -> bytes:

    pdf = LegalEasePDF(
        doc_type.strip()
        or "Legal Document"
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        size=11
    )

    for line in split_lines(text):

        if not line:

            pdf.ln(4)

            continue

        if is_heading(line):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                0,
                7,
                line
            )

            pdf.set_font(
                "Helvetica",
                size=11
            )

        elif is_bullet(line):

            bullet = re.sub(
                r"^[-*•]\s+",
                "",
                line
            )

            pdf.multi_cell(
                0,
                6,
                "- " + bullet
            )

        else:

            pdf.multi_cell(
                0,
                6,
                line
            )

        pdf.ln(1)

    clean_terms = [
        item.strip()
        for item in terms.split(";")
        if item.strip()
    ]

    if clean_terms:

        pdf.ln(4)

        pdf.set_font(
            "Helvetica",
            "B",
            12
        )

        pdf.cell(
            0,
            7,
            "Key Terms Provided"
        )

        pdf.ln(8)

        pdf.set_font(
            "Helvetica",
            size=10
        )

        for index, term in enumerate(
            clean_terms,
            start=1
        ):

            pdf.multi_cell(
                0,
                6,
                f"{index}. {term}"
            )

    return bytes(
        pdf.output()
    )


def format_html_preview(
    text: str
) -> str:

    chunks = []

    for raw_line in sanitize_text(
        text
    ).splitlines():

        line = raw_line.strip()

        if not line:

            chunks.append(
                "<div class='doc-spacer'></div>"
            )

            continue

        safe = html.escape(line)

        if is_heading(line):

            chunks.append(
                f"<h3>{safe}</h3>"
            )

        elif is_bullet(line):

            bullet = re.sub(
                r"^[-*•]\s+",
                "",
                line
            )

            chunks.append(
                "<div class='bullet'>"
                f"• {html.escape(bullet)}"
                "</div>"
            )

        else:

            chunks.append(
                f"<p>{safe}</p>"
            )

    return "\n".join(chunks)